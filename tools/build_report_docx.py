"""Builds a Word (.docx) version of the DSA521S Group 11 project report.

It reuses the exact content of tools/build_report.py by re-implementing that
script's helper API (P/H1/H2/H3/table/code/shot/bullets) as Word emitters and
executing the content section of the file.
"""
import os, re, html
from docx import Document
from docx.shared import Pt, Mm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

BASE = "/home/user/workspace/DSA521S_Project"
SHOTS = os.path.join(BASE, "screenshots")
OUT = os.path.join(BASE, "DSA521S_Group11_Project_Report.docx")
REPO_URL = open(os.path.join(BASE, "tools", "repo_url.txt")).read().strip()

INK = RGBColor(0x28, 0x25, 0x1D)
MUTED = RGBColor(0x5F, 0x5D, 0x57)
TEAL = RGBColor(0x01, 0x69, 0x6F)
TEAL_HEX = "01696F"
ZEBRA = "F9F8F5"
SOFT = "F1EFE9"
BODY_FONT = "Calibri"
MONO_FONT = "Courier New"
MONO = MONO_FONT

USABLE_MM = 165.0          # same text width as the PDF
mm = 1.0                   # content uses "165 * mm" for widths -> treat as mm

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Mm(210), Mm(297)
sec.left_margin = sec.right_margin = Mm(22.5)
sec.top_margin = Mm(20)
sec.bottom_margin = Mm(18)

st = doc.styles["Normal"]
st.font.name = BODY_FONT
st.font.size = Pt(10.5)
st.font.color.rgb = INK
st.paragraph_format.space_after = Pt(6)
st.paragraph_format.line_spacing = 1.12


def _shade(el, fill):
    sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear")
    sh.set(qn("w:fill"), fill)
    el.append(sh)


def _borders(tbl):
    tblPr = tbl._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement("w:" + edge)
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), "4")
        e.set(qn("w:color"), "D4D1CA")
        borders.append(e)
    tblPr.append(borders)


def _repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    trPr.append(el)


def _cant_split(row):
    trPr = row._tr.get_or_add_trPr()
    trPr.append(OxmlElement("w:cantSplit"))


def _hyperlink(par, url, text):
    part = par.part
    r_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), r_id)
    r = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")
    col = OxmlElement("w:color")
    col.set(qn("w:val"), TEAL_HEX)
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    rPr.append(col)
    rPr.append(u)
    r.append(rPr)
    t = OxmlElement("w:t")
    t.text = text
    r.append(t)
    link.append(r)
    par._p.append(link)


ENTITIES = {"&mdash;": "\u2014", "&ndash;": "\u2013", "&minus;": "\u2212",
            "&times;": "\u00d7", "&divide;": "\u00f7", "&rarr;": "\u2192",
            "&le;": "\u2264", "&ge;": "\u2265", "&gt;": ">", "&lt;": "<",
            "&sup2;": "\u00b2", "&amp;": "&", "&nbsp;": " "}

TAG = re.compile(r"<(/?)(b|i|font|a)([^>]*)>", re.I)


def emit_runs(par, markup, *, size=None, color=None, bold=False, font=None):
    """Render a small subset of ReportLab inline markup into Word runs."""
    text = markup
    for k, v in ENTITIES.items():
        text = text.replace(k, v)
    text = re.sub(r"&#(\d+);", lambda m: chr(int(m.group(1))), text)

    pos = 0
    stack = {"b": bold, "i": False, "mono": False, "link": None}
    for m in TAG.finditer(text):
        chunk = text[pos:m.start()]
        if chunk:
            _add_run(par, chunk, stack, size, color, font)
        closing, tag, attrs = m.group(1) == "/", m.group(2).lower(), m.group(3)
        if tag == "b":
            stack["b"] = bold if closing else True
        elif tag == "i":
            stack["i"] = not closing
        elif tag == "font":
            stack["mono"] = not closing
        elif tag == "a":
            stack["link"] = None if closing else re.search(r'href="([^"]+)"', attrs).group(1)
        pos = m.end()
    tail = text[pos:]
    if tail:
        _add_run(par, tail, stack, size, color, font)


def _add_run(par, text, stack, size, color, font):
    if stack["link"]:
        _hyperlink(par, stack["link"], text)
        return
    r = par.add_run(text)
    r.bold = stack["b"]
    r.italic = stack["i"]
    r.font.name = MONO_FONT if stack["mono"] else (font or BODY_FONT)
    if stack["mono"]:
        r.font.size = Pt(9)
    elif size:
        r.font.size = Pt(size)
    if color is not None:
        r.font.color.rgb = color


# ---------------- helper API mirroring build_report.py ----------------

SPEC = {
    "title": dict(size=19, bold=True, color=INK, align="c", after=4),
    "sub": dict(size=12, color=MUTED, align="c", after=4),
    "body": dict(size=10.5, color=INK, after=6),
    "cap": dict(size=8.5, color=MUTED, align="c", after=10, before=3),
}


def P(t, s="body"):
    cfg = SPEC[s]
    par = doc.add_paragraph()
    pf = par.paragraph_format
    pf.space_after = Pt(cfg.get("after", 6))
    pf.space_before = Pt(cfg.get("before", 0))
    if cfg.get("align") == "c":
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    emit_runs(par, t, size=cfg["size"], color=cfg.get("color"), bold=cfg.get("bold", False))
    return par


def H1(t):
    par = doc.add_paragraph()
    par.paragraph_format.space_before = Pt(16)
    par.paragraph_format.space_after = Pt(2)
    par.paragraph_format.keep_with_next = True
    emit_runs(par, t, size=15.5, color=TEAL, bold=True)
    rule = doc.add_paragraph()
    rule.paragraph_format.space_before = Pt(0)
    rule.paragraph_format.space_after = Pt(8)
    rule.paragraph_format.keep_with_next = True
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "12")
    bottom.set(qn("w:color"), TEAL_HEX)
    pBdr.append(bottom)
    rule._p.get_or_add_pPr().append(pBdr)


def H2(t):
    par = doc.add_paragraph()
    par.paragraph_format.space_before = Pt(12)
    par.paragraph_format.space_after = Pt(4)
    par.paragraph_format.keep_with_next = True
    emit_runs(par, t, size=12, color=INK, bold=True)


def H3(t):
    par = doc.add_paragraph()
    par.paragraph_format.space_before = Pt(9)
    par.paragraph_format.space_after = Pt(3)
    par.paragraph_format.keep_with_next = True
    emit_runs(par, t, size=10.5, color=MUTED, bold=True)


def bullets(items):
    for it in items:
        par = doc.add_paragraph(style="List Bullet")
        par.paragraph_format.space_after = Pt(4)
        emit_runs(par, it, size=10.5, color=INK)


def code(text, size=8.2):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.LEFT
    _borders(tbl)
    cell = tbl.cell(0, 0)
    cell.width = Mm(USABLE_MM)
    _shade(cell._tc.get_or_add_tcPr(), SOFT)
    lines = text.strip("\n").split("\n")
    for i, line in enumerate(lines):
        par = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
        pf = par.paragraph_format
        pf.space_after = Pt(0)
        pf.space_before = Pt(0)
        pf.line_spacing = 1.0
        r = par.add_run(line.replace(" ", "\u00a0"))
        r.font.name = MONO_FONT
        r.font.size = Pt(size)
        r.font.color.rgb = INK
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(4)
    spacer.paragraph_format.space_before = Pt(0)


def table(rows, widths, align_right=()):
    ncols = len(rows[0])
    tbl = doc.add_table(rows=0, cols=ncols)
    tbl.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl.autofit = False
    _borders(tbl)
    total = sum(widths)
    scaled = [w / total * USABLE_MM for w in widths]
    for i, raw in enumerate(rows):
        row = tbl.add_row()
        _cant_split(row)
        if i == 0:
            _repeat_header(row)
        for j, val in enumerate(raw):
            cell = row.cells[j]
            cell.width = Mm(scaled[j])
            tcPr = cell._tc.get_or_add_tcPr()
            if i == 0:
                _shade(tcPr, TEAL_HEX)
            elif i % 2 == 0:
                _shade(tcPr, ZEBRA)
            par = cell.paragraphs[0]
            par.paragraph_format.space_after = Pt(0)
            par.paragraph_format.line_spacing = 1.08
            if j in align_right:
                par.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            emit_runs(par, str(val), size=9.5,
                      color=RGBColor(0xFF, 0xFF, 0xFF) if i == 0 else INK,
                      bold=(i == 0))
    for j, w in enumerate(scaled):
        for row in tbl.rows:
            row.cells[j].width = Mm(w)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(6)
    spacer.paragraph_format.space_before = Pt(0)


def shot(name, caption, max_w=165.0, max_h=185.0):
    from PIL import Image as PILImage
    path = os.path.join(SHOTS, name)
    iw, ih = PILImage.open(path).size
    px_per_mm = iw / max_w
    height_mm = ih / px_per_mm
    if height_mm > max_h:
        max_w = max_w * max_h / height_mm
    par = doc.add_paragraph()
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    par.paragraph_format.space_after = Pt(2)
    par.paragraph_format.keep_with_next = True
    par.add_run().add_picture(path, width=Mm(max_w))
    P(caption, "cap")


# Stubs for the layout primitives used by the shared content section.
class Spacer:
    def __init__(self, w, h):
        self.h = h


class PageBreak:
    pass


def A(item):
    if isinstance(item, PageBreak):
        doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    elif isinstance(item, Spacer):
        par = doc.add_paragraph()
        par.paragraph_format.space_after = Pt(max(2.0, item.h * 2.0))
        par.paragraph_format.space_before = Pt(0)
        for r in par.runs:
            r.font.size = Pt(2)


# ---------------- footer ----------------
footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.LEFT
fr = footer.add_run("DSA521S Group Mini-Project 2026 \u2014 Group 11 \u2014 NUST Service Centre Simulation\t\tPage ")
fr.font.size = Pt(8)
fr.font.color.rgb = MUTED
fld = OxmlElement("w:fldSimple")
fld.set(qn("w:instr"), "PAGE")
run = OxmlElement("w:r")
rPr = OxmlElement("w:rPr")
szEl = OxmlElement("w:sz")
szEl.set(qn("w:val"), "16")
rPr.append(szEl)
run.append(rPr)
fld.append(run)
footer._p.append(fld)


# ---------------- run the shared content ----------------
src = open(os.path.join(BASE, "tools", "build_report.py")).read()
marker = "# ======================= COVER ======================="
content = src[src.index(marker):]
content = content[:content.index("# ======================= BUILD =======================")]
exec(compile(content, "report_content", "exec"), globals())

doc.core_properties.author = "Perplexity Computer"
doc.core_properties.title = "DSA521S Group 11 Mini-Project 2026 — Project Report"
doc.save(OUT)
print("wrote", OUT)
