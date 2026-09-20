"""Converts the project's Markdown documents (README, PSEUDOCODE, GITHUB_SETUP)
into Word documents styled like the project report."""
import os, re, sys
from docx import Document
from docx.shared import Pt, Mm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import sys; sys.path.insert(0, "/home/user/workspace/DSA521S_Project/tools")
from ooxml_order import add_tbl_pr, add_p_pr

BASE = "/home/user/workspace/DSA521S_Project"
INK = RGBColor(0x28, 0x25, 0x1D)
MUTED = RGBColor(0x5F, 0x5D, 0x57)
TEAL = RGBColor(0x01, 0x69, 0x6F)
TEAL_HEX = "01696F"
SOFT = "F1EFE9"
ZEBRA = "F9F8F5"
BODY_FONT = "Calibri"
MONO = "Courier New"


def shade(el, fill):
    sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear")
    sh.set(qn("w:fill"), fill)
    el.append(sh)


def borders(tbl):
    b = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement("w:" + edge)
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), "4")
        e.set(qn("w:color"), "D4D1CA")
        b.append(e)
    add_tbl_pr(tbl._tbl.tblPr, b)


def hyperlink(par, url, text):
    r_id = par.part.relate_to(
        url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), r_id)
    r = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")
    c = OxmlElement("w:color"); c.set(qn("w:val"), TEAL_HEX)
    u = OxmlElement("w:u"); u.set(qn("w:val"), "single")
    rPr.append(c); rPr.append(u); r.append(rPr)
    t = OxmlElement("w:t"); t.text = text; t.set(qn("xml:space"), "preserve")
    r.append(t); link.append(r); par._p.append(link)


TOKEN = re.compile(r"(\[[^\]]+\]\([^)]+\)|\*\*[^*]+\*\*|`[^`]+`|https?://\S+)")


def inline(par, text, size=10.5, bold=False, color=INK, mono_size=9.5):
    text = text.replace("\\|", "|")
    text = re.sub(r"<(https?://[^>]+)>", r"\1", text)
    for part in TOKEN.split(text):
        if not part:
            continue
        m = re.fullmatch(r"\[([^\]]+)\]\(([^)]+)\)", part)
        if m:
            hyperlink(par, m.group(2), m.group(1))
            continue
        if re.fullmatch(r"https?://\S+", part):
            hyperlink(par, part.rstrip(".,)"), part.rstrip(".,)"))
            continue
        if part.startswith("**") and part.endswith("**"):
            r = par.add_run(part[2:-2]); r.bold = True
            r.font.size = Pt(size); r.font.name = BODY_FONT; r.font.color.rgb = color
            continue
        if part.startswith("`") and part.endswith("`"):
            r = par.add_run(part[1:-1]); r.font.name = MONO
            r.font.size = Pt(mono_size); r.font.color.rgb = color; r.bold = bold
            continue
        r = par.add_run(part)
        r.bold = bold
        r.font.size = Pt(size); r.font.name = BODY_FONT; r.font.color.rgb = color


def convert(md_path, out_path, title, landscape_code=False, code_size=8.5):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Mm(210), Mm(297)
    sec.left_margin = sec.right_margin = Mm(20)
    sec.top_margin = Mm(20); sec.bottom_margin = Mm(18)
    usable = 170.0

    n = doc.styles["Normal"]
    n.font.name = BODY_FONT
    n.font.size = Pt(10.5)
    n.font.color.rgb = INK
    n.paragraph_format.space_after = Pt(6)
    n.paragraph_format.line_spacing = 1.12

    f = sec.footer.paragraphs[0]
    fr = f.add_run("DSA521S Group Mini-Project 2026 \u2014 Group 11 \u2014 %s\t\tPage " % title)
    fr.font.size = Pt(8); fr.font.color.rgb = MUTED
    fld = OxmlElement("w:fldSimple"); fld.set(qn("w:instr"), "PAGE")
    r = OxmlElement("w:r"); rPr = OxmlElement("w:rPr")
    sz = OxmlElement("w:sz"); sz.set(qn("w:val"), "16"); rPr.append(sz); r.append(rPr)
    fld.append(r); f._p.append(fld)

    lines = open(md_path).read().split("\n")
    i = 0
    while i < len(lines):
        line = lines[i]
        s = line.strip()

        if s.startswith("```"):
            i += 1
            buf = []
            while i < len(lines) and not lines[i].strip().startswith("```"):
                buf.append(lines[i]); i += 1
            i += 1
            tbl = doc.add_table(rows=1, cols=1)
            tbl.alignment = WD_TABLE_ALIGNMENT.LEFT
            borders(tbl)
            cell = tbl.cell(0, 0)
            cell.width = Mm(usable)
            shade(cell._tc.get_or_add_tcPr(), SOFT)
            for k, cl in enumerate(buf):
                p = cell.paragraphs[0] if k == 0 else cell.add_paragraph()
                pf = p.paragraph_format
                pf.space_after = Pt(0); pf.space_before = Pt(0); pf.line_spacing = 1.0
                pf.left_indent = Mm(8); pf.first_line_indent = Mm(-8)
                rr = p.add_run(cl.replace("\t", "    ").replace(" ", "\u00a0"))
                rr.font.name = MONO; rr.font.size = Pt(code_size); rr.font.color.rgb = INK
            sp = doc.add_paragraph()
            sp.paragraph_format.space_after = Pt(6); sp.paragraph_format.space_before = Pt(0)
            continue

        if s.startswith("|") and i + 1 < len(lines) and re.match(r"^\|[\s:\-|]+\|$", lines[i + 1].strip()):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                raw = lines[i].strip()
                if not re.match(r"^\|[\s:\-|]+\|$", raw):
                    rows.append([c.strip() for c in raw.strip("|").split("|")])
                i += 1
            ncols = max(len(r) for r in rows)
            tbl = doc.add_table(rows=0, cols=ncols)
            tbl.alignment = WD_TABLE_ALIGNMENT.LEFT
            tbl.autofit = False
            borders(tbl)
            weights = [max(len(r[c]) if c < len(r) else 0 for r in rows) for c in range(ncols)]
            tot = sum(weights) or 1
            widths = [max(18.0, w / tot * usable) for w in weights]
            scale = usable / sum(widths)
            widths = [w * scale for w in widths]
            for ri, raw in enumerate(rows):
                row = tbl.add_row()
                trPr = row._tr.get_or_add_trPr()
                trPr.append(OxmlElement("w:cantSplit"))
                if ri == 0:
                    h = OxmlElement("w:tblHeader"); h.set(qn("w:val"), "true"); trPr.append(h)
                for ci in range(ncols):
                    cell = row.cells[ci]
                    cell.width = Mm(widths[ci])
                    tcPr = cell._tc.get_or_add_tcPr()
                    if ri == 0:
                        shade(tcPr, TEAL_HEX)
                    elif ri % 2 == 0:
                        shade(tcPr, ZEBRA)
                    p = cell.paragraphs[0]
                    p.paragraph_format.space_after = Pt(0)
                    p.paragraph_format.line_spacing = 1.08
                    inline(p, raw[ci] if ci < len(raw) else "", size=9.5,
                           bold=(ri == 0),
                           color=RGBColor(0xFF, 0xFF, 0xFF) if ri == 0 else INK,
                           mono_size=8.8)
            sp = doc.add_paragraph()
            sp.paragraph_format.space_after = Pt(8); sp.paragraph_format.space_before = Pt(0)
            continue

        if re.match(r"^#{1,4} ", s):
            level = len(s) - len(s.lstrip("#"))
            text = s[level:].strip()
            p = doc.add_paragraph()
            pf = p.paragraph_format
            pf.keep_with_next = True
            if level == 1:
                pf.space_before = Pt(0); pf.space_after = Pt(6)
                inline(p, text, size=18, bold=True, color=INK)
            elif level == 2:
                pf.space_before = Pt(16); pf.space_after = Pt(2)
                inline(p, text, size=14.5, bold=True, color=TEAL)
                rp = doc.add_paragraph()
                rp.paragraph_format.space_before = Pt(0)
                rp.paragraph_format.space_after = Pt(8)
                rp.paragraph_format.keep_with_next = True
                pBdr = OxmlElement("w:pBdr")
                b = OxmlElement("w:bottom")
                b.set(qn("w:val"), "single"); b.set(qn("w:sz"), "12"); b.set(qn("w:color"), TEAL_HEX)
                pBdr.append(b)
                add_p_pr(p._p.get_or_add_pPr(), pBdr)
            else:
                pf.space_before = Pt(11); pf.space_after = Pt(3)
                inline(p, text, size=11.5, bold=True, color=INK)
            i += 1
            continue

        if re.match(r"^(\*|-) ", s) or re.match(r"^\d+\. ", s):
            ordered = bool(re.match(r"^\d+\. ", s))
            indent = len(line) - len(line.lstrip())
            style = "List Number" if ordered else "List Bullet"
            p = doc.add_paragraph(style=style)
            p.paragraph_format.space_after = Pt(4)
            if indent >= 2:
                p.paragraph_format.left_indent = Mm(16)
            body = [re.sub(r"^(\*|-|\d+\.)\s+", "", s)]
            i += 1
            while i < len(lines) and lines[i].strip() and not re.match(
                    r"^(#{1,4} |```|\||\*\s|-\s|\d+\.\s|---)", lines[i].strip()):
                body.append(lines[i].strip())
                i += 1
            inline(p, " ".join(body))
            continue

        if s.startswith("---") or s == "":
            i += 1
            continue

        buf = [s]
        i += 1
        while i < len(lines) and lines[i].strip() and not re.match(
                r"^(#{1,4} |```|\||\*|-|\d+\. |---)", lines[i].strip()):
            buf.append(lines[i].strip()); i += 1
        p = doc.add_paragraph()
        inline(p, " ".join(buf))

    doc.core_properties.author = "Perplexity Computer"
    doc.core_properties.title = "DSA521S Group 11 — " + title
    doc.save(out_path)
    print("wrote", out_path)


JOBS = [
    ("README.md", "DSA521S_Group11_README.docx", "README", 8.0),
    ("PSEUDOCODE.md", "DSA521S_Group11_PSEUDOCODE.docx", "Part E pseudocode", 9.0),
    ("GITHUB_SETUP.md", "DSA521S_Group11_GITHUB_SETUP.docx", "GitHub setup guide", 7.0),
]

for src, out, title, csize in JOBS:
    convert(os.path.join(BASE, src), os.path.join(BASE, out), title, code_size=csize)
