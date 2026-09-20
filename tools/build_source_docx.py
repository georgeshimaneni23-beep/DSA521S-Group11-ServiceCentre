"""Builds a Word document containing all Java source files of the project."""
import os
from docx import Document
from docx.shared import Pt, Mm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import sys; sys.path.insert(0, "/home/user/workspace/DSA521S_Project/tools")
from ooxml_order import add_tbl_pr, add_p_pr
from docx.enum.section import WD_ORIENT

BASE = "/home/user/workspace/DSA521S_Project"
SRC = os.path.join(BASE, "src")
OUT = os.path.join(BASE, "DSA521S_Group11_Source_Code.docx")

ORDER = ["Student.java", "StudentQueue.java", "ServiceRecordList.java", "IntStack.java",
         "PostfixEvaluator.java", "DailyStatistics.java", "ArrayUtil.java", "Sorters.java",
         "SortingExperiment.java", "ServiceCentreApp.java", "DemoRunner.java"]

PURPOSE = {
    "Student.java": "Part A — shared student record: student number, name, service type and estimated service time.",
    "StudentQueue.java": "Task A1 — the waiting line, a FIFO queue of linked nodes with front and rear pointers.",
    "ServiceRecordList.java": "Task A2 — singly linked list of service records: insert, delete, search and traverse.",
    "IntStack.java": "Task A3 — array-based LIFO stack of integers used by the postfix evaluator.",
    "PostfixEvaluator.java": "Task A3 — evaluates postfix expressions using the stack, with a step-by-step trace.",
    "DailyStatistics.java": "Task A4 — array traversal for total, average, highest, lowest and threshold counts.",
    "ArrayUtil.java": "Helper — array copying, printing and generation used by Parts B and C.",
    "Sorters.java": "Part B — Selection, Insertion, Merge and Quick Sort, all implemented from first principles.",
    "SortingExperiment.java": "Part C — counts comparisons and measures times for the four sorts at several input sizes.",
    "ServiceCentreApp.java": "Part D — the integrated menu-driven system that drives every structure above.",
    "DemoRunner.java": "Evidence — prints every demonstration and trace reproduced in the project report.",
}

INK = RGBColor(0x28, 0x25, 0x1D)
MUTED = RGBColor(0x5F, 0x5D, 0x57)
TEAL = RGBColor(0x01, 0x69, 0x6F)

doc = Document()
sec = doc.sections[0]
sec.orientation = WD_ORIENT.LANDSCAPE
sec.page_width, sec.page_height = Mm(297), Mm(210)
sec.left_margin = sec.right_margin = Mm(18)
sec.top_margin = Mm(18)
sec.bottom_margin = Mm(16)

n = doc.styles["Normal"]
n.font.name = "Calibri"
n.font.size = Pt(10.5)
n.font.color.rgb = INK

footer = sec.footer.paragraphs[0]
fr = footer.add_run("DSA521S Group Mini-Project 2026 — Group 11 — Java source code\t\tPage ")
fr.font.size = Pt(8)
fr.font.color.rgb = MUTED
fld = OxmlElement("w:fldSimple")
fld.set(qn("w:instr"), "PAGE")
r = OxmlElement("w:r")
rPr = OxmlElement("w:rPr")
sz = OxmlElement("w:sz")
sz.set(qn("w:val"), "16")
rPr.append(sz)
r.append(rPr)
fld.append(r)
footer._p.append(fld)


def para(text, size=10.5, bold=False, color=INK, align=None, before=0, after=6, keep=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.keep_with_next = keep
    if align == "c":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    run.font.color.rgb = color
    return p


def rule(colour="01696F"):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.keep_with_next = True
    pBdr = OxmlElement("w:pBdr")
    b = OxmlElement("w:bottom")
    b.set(qn("w:val"), "single")
    b.set(qn("w:sz"), "12")
    b.set(qn("w:color"), colour)
    pBdr.append(b)
    add_p_pr(p._p.get_or_add_pPr(), pBdr)


# ---------------- cover ----------------
para("NAMIBIA UNIVERSITY OF SCIENCE AND TECHNOLOGY", 12, color=MUTED, align="c", before=30, after=2)
para("DSA521S — DATA STRUCTURES AND ALGORITHMS 1", 12, color=MUTED, align="c", after=24)
para("NUST Service Centre Simulation", 19, bold=True, align="c", after=4)
para("Complete Java Source Code Listing", 19, bold=True, align="c", after=18)
para("Group Mini-Project 2026 — Group 11", 12, color=MUTED, align="c", after=40)
para("Submitted by: 225018497 — George Shimaneni", 11, color=MUTED, align="c", after=4)
para("George Shimaneni (225018497) · Dipundhi Paul Peter (225058146) · "
     "Mangulukeni Kayoko (227074404) · Nelumbu Rachel (225049511) · Enerist T Shilumbu (222093951)",
     10, color=MUTED, align="c", after=24)
para("This listing contains the 11 Java files of the project in the order in which they are "
     "discussed in the project report. Compile with  javac -d out src/*.java  and run with "
     "java -cp out ServiceCentreApp.", 10, color=MUTED, align="c")
doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

# ---------------- contents ----------------
para("Contents", 15, bold=True, color=TEAL, after=2)
rule()
for i, name in enumerate(ORDER, 1):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    r1 = p.add_run("%d. %s — " % (i, name))
    r1.bold = True
    r1.font.size = Pt(10.5)
    r2 = p.add_run(PURPOSE[name])
    r2.font.size = Pt(10.5)
    r2.font.color.rgb = MUTED

# ---------------- listings ----------------
for i, name in enumerate(ORDER, 1):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    para("%d. %s" % (i, name), 15, bold=True, color=TEAL, after=2, keep=True)
    rule()
    para(PURPOSE[name], 10.5, color=MUTED, after=10, keep=True)
    lines = open(os.path.join(SRC, name)).read().rstrip("\n").split("\n")
    width = len(str(len(lines)))
    for ln, line in enumerate(lines, 1):
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.space_before = Pt(0)
        pf.space_after = Pt(0)
        pf.line_spacing = 1.0
        pf.left_indent = Mm(10)
        pf.first_line_indent = Mm(-10)
        num = p.add_run(str(ln).rjust(width) + "  ")
        num.font.name = "Courier New"
        num.font.size = Pt(8)
        num.font.color.rgb = RGBColor(0xA8, 0xA6, 0xA0)
        code = p.add_run(line.replace("\t", "    ").replace(" ", "\u00a0"))
        code.font.name = "Courier New"
        code.font.size = Pt(8.5)
        code.font.color.rgb = INK

doc.core_properties.author = "Perplexity Computer"
doc.core_properties.title = "DSA521S Group 11 — Java Source Code Listing"
doc.save(OUT)
print("wrote", OUT)
